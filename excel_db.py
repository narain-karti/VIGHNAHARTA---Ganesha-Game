import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_FILE = "vighnaharta_scores.xlsx"

CAMPUSES = [
    # Telangana / Hyderabad
    "NIAT - Malla Reddy Vishwavidyapeeth (MRV), Hyderabad",
    "NIAT - Aurora Deemed University, Hyderabad",
    "NIAT - Chaitanya Deemed University, Hyderabad",
    "NIAT - Kapil Kavuri Hub, Hyderabad",
    # Karnataka / Bengaluru
    "NIAT - S-VYASA School of Advanced Studies, Bengaluru",
    "NIAT - St. Peter's Institute (SPIHER), Bengaluru",
    "NIAT - Yenepoya University, Mangalore",
    # Andhra Pradesh
    "NIAT - Chalapathi Institute of Technology (CIT), Guntur",
    "NIAT - Lingaya's Institute (LIMAT), Vijayawada",
    "NIAT - NRI University, Vijayawada",
    "NIAT - NSRIT, Visakhapatnam",
    "NIAT - BEST Innovation University, Anantapur",
    "NIAT - Annamacharya University, Kadapa",
    # Tamil Nadu & Pondicherry
    "NIAT - Takshashila University, Chennai-Pondicherry",
    "NIAT - AMET University, Chennai",
    "NIAT - B.S. Abdur Rahman Crescent Institute, Chennai",
    "NIAT - Joy University, Tirunelveli",
    # Maharashtra
    "NIAT - Ajeenkya DY Patil University, Pune",
    "NIAT - Sanjay Ghodawat University, Kolhapur",
    # North & Central India
    "NIAT - Noida International University (NIU), Greater Noida",
    "NIAT - Sanskriti University, Mathura",
    "NIAT - Vivekananda Global University (VGU), Jaipur",
    "NIAT - Rabindranath Tagore University, Bhopal",
    # Other
    "NIAT - Other / Online COE"
]

def create_initial_workbook():
    wb = openpyxl.Workbook()
    
    # Sheet 1: Devotee Game Sessions
    ws1 = wb.active
    ws1.title = "Devotee Scores"
    
    headers_ws1 = [
        "Session ID", "Timestamp", "Devotee Name", "NIAT Campus", 
        "Student ID", "Punya Score", "Wave Reached", "Obstacles Cleared", 
        "Best Combo", "Diyas Left", "Status"
    ]
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="B8860B", end_color="B8860B", fill_type="solid") # Dark Goldenrod
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    ws1.append(headers_ws1)
    for col_num in range(1, len(headers_ws1) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Sheet 2: Campus Leaderboard
    ws2 = wb.create_sheet(title="Campus Leaderboard")
    headers_ws2 = [
        "Rank", "NIAT Campus", "Total Devotees", "Total Combined Score", 
        "Highest Single Score", "Top Devotee"
    ]
    ws2.append(headers_ws2)
    header_fill_campus = PatternFill(start_color="D2691E", end_color="D2691E", fill_type="solid") # Chocolate / Saffron
    for col_num in range(1, len(headers_ws2) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill_campus
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for ws in (ws1, ws2):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
            
    wb.save(EXCEL_FILE)
    print(f"Created initial Excel database: {EXCEL_FILE}")

def append_score_record(name, campus, score, wave, cleared, combo, student_id="", blessings=0):
    if not os.path.exists(EXCEL_FILE):
        create_initial_workbook()
        
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws1 = wb["Devotee Scores"]
    
    row_count = ws1.max_row
    session_id = f"VGH-{datetime.now().strftime('%Y%m%d')}-{row_count:04d}"
    timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    status = "Sacred Record" if score >= 1000 else "Puja Completed"
    
    new_row = [
        session_id, timestamp_str, name, campus,
        student_id or "N/A", int(score), int(wave), int(cleared),
        f"x{combo}", int(blessings), status
    ]
    ws1.append(new_row)
    
    # Border & formatting for the new row
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    current_row = ws1.max_row
    for col_num in range(1, len(new_row) + 1):
        cell = ws1.cell(row=current_row, column=col_num)
        cell.border = thin_border
        if col_num in (6, 7, 8, 10): # Numbers
            cell.alignment = Alignment(horizontal="right")
        elif col_num in (1, 2, 5, 9, 11):
            cell.alignment = Alignment(horizontal="center")
            
    # Update Sheet 2: Campus Leaderboard
    ws2 = wb["Campus Leaderboard"]
    # Clear old data rows in ws2
    while ws2.max_row > 1:
        ws2.delete_rows(2)
        
    # Aggregate data from ws1
    campus_stats = {}
    for r in range(2, ws1.max_row + 1):
        c_name = ws1.cell(row=r, column=4).value
        p_name = ws1.cell(row=r, column=3).value
        s_val = ws1.cell(row=r, column=6).value or 0
        if not c_name:
            continue
        if c_name not in campus_stats:
            campus_stats[c_name] = {
                "devotees": set(),
                "total_score": 0,
                "highest_score": 0,
                "top_devotee": p_name
            }
        campus_stats[c_name]["devotees"].add(p_name)
        campus_stats[c_name]["total_score"] += int(s_val)
        if int(s_val) > campus_stats[c_name]["highest_score"]:
            campus_stats[c_name]["highest_score"] = int(s_val)
            campus_stats[c_name]["top_devotee"] = p_name

    # Sort campuses by total score descending
    sorted_campuses = sorted(campus_stats.items(), key=lambda x: x[1]["total_score"], reverse=True)
    for rank, (c_name, stats) in enumerate(sorted_campuses, 1):
        ws2.append([
            rank, c_name, len(stats["devotees"]), stats["total_score"],
            stats["highest_score"], stats["top_devotee"]
        ])
        r_idx = ws2.max_row
        for c_idx in range(1, 7):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if c_idx in (1, 3):
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in (4, 5):
                cell.alignment = Alignment(horizontal="right")

    # Refresh column widths
    for ws in (ws1, ws2):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(EXCEL_FILE)
    print(f"Recorded score for {name} ({campus}): {score} pts into {EXCEL_FILE}")
    return session_id

if __name__ == "__main__":
    create_initial_workbook()
    # Add a sample record for demonstration
    append_score_record("Karthik Narain", "NIAT - Malla Reddy Vishwavidyapeeth (MRV), Hyderabad", 1250, 4, 38, 4, "NIAT-2024-001", 3)
